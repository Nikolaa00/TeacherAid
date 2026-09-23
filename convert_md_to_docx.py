import re
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_document():
    doc = docx.Document()
    
    # Page setup - Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    
    with open('/workspace/questions.md', 'r', encoding='utf-8') as f:
        md_text = f.read()

    lines = md_text.split('\n')
    
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        
        # Blank line
        if not stripped:
            i += 1
            continue
            
        # Title (# )
        if line.startswith('# '):
            text = line[2:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(10)
            run = p.add_run(text)
            run.font.size = Pt(20)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D) # Navy
            i += 1
            continue
            
        # Heading 2 (## )
        if line.startswith('## '):
            text = line[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(text)
            run.font.size = Pt(15)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x2B, 0x54, 0x7E) # Slate Blue
            # Bottom border / divider line under H2
            pPr = p._p.get_or_add_pPr()
            pbdr = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                             r'<w:bottom w:val="single" w:sz="12" w:space="4" w:color="2B547E"/>'
                             r'</w:pBdr>')
            pPr.append(pbdr)
            i += 1
            continue
            
        # Heading 3 (### )
        if line.startswith('### '):
            text = line[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(text)
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x66)
            i += 1
            continue
            
        # Heading 4 (#### )
        if line.startswith('#### '):
            text = line[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(text)
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x44, 0x44, 0x77)
            i += 1
            continue
            
        # Heading 5 (##### )
        if line.startswith('##### '):
            text = line[6:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(text)
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.italic = True
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
            i += 1
            continue
            
        # Horizontal Rule (---)
        if stripped == '---':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            pPr = p._p.get_or_add_pPr()
            pbdr = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                             r'<w:bottom w:val="single" w:sz="6" w:space="1" w:color="CCCCCC"/>'
                             r'</w:pBdr>')
            pPr.append(pbdr)
            i += 1
            continue
            
        # Math Block ($$ ... $$)
        if stripped.startswith('$$') and stripped.endswith('$$') and len(stripped) > 4:
            math_text = stripped[2:-2].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            # Add with light background or subtle box
            pPr = p._p.get_or_add_pPr()
            shd = parse_xml(r'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="clear" w:color="auto" w:fill="F5F5F5"/>')
            pPr.append(shd)
            run = p.add_run(math_text)
            run.font.name = 'Cambria Math'
            run.font.size = Pt(11.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
            i += 1
            continue
            
        # Multi-line Math Block
        if stripped == '$$':
            math_lines = []
            i += 1
            while i < len(lines) and lines[i].strip() != '$$':
                math_lines.append(lines[i].strip())
                i += 1
            i += 1 # skip closing $$
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            pPr = p._p.get_or_add_pPr()
            shd = parse_xml(r'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="clear" w:color="auto" w:fill="F5F5F5"/>')
            pPr.append(shd)
            for m_line in math_lines:
                run = p.add_run(m_line + '\n')
                run.font.name = 'Cambria Math'
                run.font.size = Pt(11)
            continue

        # Bullet or numbered item
        indent_level = 0
        bullet_match = re.match(r'^(\s*)([-*]|\d+\.)\s+(.*)$', line)
        
        if bullet_match:
            whitespace, bullet_type, content = bullet_match.groups()
            indent_spaces = len(whitespace)
            
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            
            # Left indentation based on spaces
            if indent_spaces == 0:
                p.paragraph_format.left_indent = Inches(0.25)
            elif indent_spaces <= 2:
                p.paragraph_format.left_indent = Inches(0.45)
            elif indent_spaces <= 4:
                p.paragraph_format.left_indent = Inches(0.65)
            else:
                p.paragraph_format.left_indent = Inches(0.85)
                
            # Bullet marker
            if bullet_type in ['-', '*']:
                marker = '•  '
                if indent_spaces >= 2:
                    marker = '–  '
            else:
                marker = bullet_type + '  '
                
            run_marker = p.add_run(marker)
            run_marker.font.bold = True
            run_marker.font.color.rgb = RGBColor(0x2B, 0x54, 0x7E)
            
            add_formatted_text(p, content)
            i += 1
            continue

        # Regular Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(3)
        add_formatted_text(p, line)
        i += 1

    doc.save('/workspace/ispitni_prasanja.docx')
    print("Successfully generated /workspace/ispitni_prasanja.docx")

def add_formatted_text(paragraph, text):
    # Regex to handle markdown formatting: bold (**), italic (*), code (`), math ($)
    # A simple tokenizer:
    pattern = re.compile(r'(\*\*\*.*?\*\*\*|\*\*.*?\*\*|\*.*?\*|`.*?`|\$\$.*?\$\$|\$.*?\$|__.*?__|[^*`$_\\]+|\\.)')
    tokens = pattern.findall(text)
    
    for token in tokens:
        if token.startswith('\\'):
            paragraph.add_run(token[1:])
        elif token.startswith('***') and token.endswith('***') and len(token) >= 6:
            r = paragraph.add_run(token[3:-3])
            r.font.bold = True
            r.font.italic = True
        elif token.startswith('**') and token.endswith('**') and len(token) >= 4:
            r = paragraph.add_run(token[2:-2])
            r.font.bold = True
        elif token.startswith('__') and token.endswith('__') and len(token) >= 4:
            r = paragraph.add_run(token[2:-2])
            r.font.bold = True
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            r = paragraph.add_run(token[1:-1])
            r.font.italic = True
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            r = paragraph.add_run(token[1:-1])
            r.font.name = 'Consolas'
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(0xA0, 0x20, 0x20)
        elif token.startswith('$') and token.endswith('$') and len(token) >= 2:
            math_str = token.strip('$')
            # Clean common LaTeX for clean reading in Word
            math_clean = clean_latex(math_str)
            r = paragraph.add_run(math_clean)
            r.font.name = 'Cambria Math'
            r.font.italic = True
        else:
            paragraph.add_run(token)

def clean_latex(latex_str):
    s = latex_str
    s = s.replace(r'\text{', '').replace(r'}', '')
    s = s.replace(r'\approx', '≈')
    s = s.replace(r'\cdot', ' · ')
    s = s.replace(r'\times', '×')
    s = s.replace(r'\propto', '∝')
    s = s.replace(r'\mu s', 'µs')
    s = s.replace(r'\Delta', 'Δ')
    s = s.replace(r'\pi', 'π')
    s = s.replace(r'\Omega', 'Ω')
    s = s.replace(r'\implies', '⟹')
    s = s.replace(r'\quad', ' ')
    s = s.replace(r'\le', '≤').replace(r'\ge', '≥').replace(r'\neq', '≠')
    s = re.sub(r'\\frac\{([^{}]+)\}\{([^{}]+)\}', r'(\1 / \2)', s)
    s = re.sub(r'\\log_\{?10\}?', 'log₁₀', s)
    s = re.sub(r'\\log_\{?2\}?', 'log₂', s)
    s = s.replace('\\', '')
    return s

if __name__ == '__main__':
    create_document()
